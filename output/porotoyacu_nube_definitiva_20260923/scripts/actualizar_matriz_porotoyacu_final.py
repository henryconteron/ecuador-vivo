"""Create a new reviewer-action matrix synchronized to the definitive cloud."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def set_cell(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    doc = Document(args.source)
    doc.paragraphs[1].runs[0].text = (
        "Porotoyacu Fault — actualización con nube definitiva, 23 September 2026"
    )
    doc.paragraphs[3].runs[0].text = doc.paragraphs[3].runs[0].text.replace(
        "La versión v4", "La versión v5"
    )
    doc.paragraphs[20].runs[0].text = (
        "Nube definitiva: NUBE_COMPLETA_FINAL_TERRENO_MICRORELIEVE_DEPURADO_UTM18S.laz "
        "(SHA-256 69106c3d56a8819a4603118a7042b082b6331894de0489e7cda7f2e31d5439b2). "
        "Nuevos controles: densidad_2m; datos/control_soporte_32_perfiles.csv; "
        "Figura_2a, Figura_2b y Figura_S1 del paquete porotoyacu_nube_definitiva_20260923. "
        "Los comentarios originales de revisores/coautores se conservan en los PDF de origen."
    )
    table = doc.tables[0]
    set_cell(table.cell(0, 1), "Cambio v5")
    set_cell(table.cell(1, 1),
             "Conteos verificados en el LAZ definitivo: 13,619,241 puntos; "
             "2,810,862 Ground (20.64%).")
    set_cell(table.cell(1, 2),
             "Figura 5: densidad, retención y soporte directo en celdas de "
             "2 × 2 m; 47.13% de celdas ocupadas tienen Ground.")
    set_cell(table.cell(2, 2),
             "52.87% de celdas ocupadas carecen de Ground directo; no hay "
             "GCP/checkpoints independientes.")
    set_cell(table.cell(3, 2),
             "12/32 perfiles A/B: P1, P2, P3, P5, P6, P7, P9, P10, P11, P18, "
             "P23 y P24. Figura 2 y control_soporte_32_perfiles.csv.")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
