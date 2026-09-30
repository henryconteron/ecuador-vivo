"""Update the photogrammetry text for the verified Agisoft context DEM.

The existing manuscript and its highlighting are preserved. Only three
photogrammetry runs are changed; profile-source descriptions stay at 1 m.
"""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1] / "manuscrito"
SOURCE = ROOT / "Manuscript_Conteron_etal_2026_FOTOGRAMETRIA_432_VERIFICADA.docx"
OUTPUT = ROOT / "Manuscript_Conteron_etal_2026_FOTOGRAMETRIA_MDT_0197m.docx"


def replace_once(run, before: str, after: str) -> None:
    if run.text.count(before) != 1:
        raise ValueError(f"Expected exactly one occurrence: {before[:60]}")
    run.text = run.text.replace(before, after)


doc = Document(SOURCE)

replace_once(
    doc.paragraphs[41].runs[0],
    "The interpolated terrain display was not treated as direct topographic measurement",
    "The extrapolated terrain display was not treated as direct topographic measurement",
)

replace_once(
    doc.paragraphs[46].runs[1],
    "(a) A 1 m terrain display based on class-2 Ground cell medians. "
    "Gaps are interpolated only for visual context, not interpreted as direct terrain "
    "observations; the observed-cell raster retains NoData.",
    "(a) A 0.196717 m/pixel terrain display generated in Agisoft Metashape from "
    "class-2 Ground points in the definitive LAZ with the Extrapolated setting. "
    "The surface is used only for visual context; extrapolated cells are not "
    "direct terrain observations. The separate 1 m observed-cell raster retains "
    "NoData where Ground points are absent.",
)

replace_once(
    doc.paragraphs[50].runs[0],
    "The final LAZ, not the initial Metashape classification, underpins the profile "
    "and density analyses. From this same LAZ, the observed 1 m terrain raster",
    "The final LAZ, not the initial Metashape classification, underpins the profile "
    "and density analyses. For the Figure 2a map only, a separate Ground-only "
    "terrain raster was generated in Agisoft Metashape using the Extrapolated "
    "setting and exported at a measured cell size of 0.196717 m. Its continuous "
    "coverage is visualization context, not direct Ground support or a measure of "
    "vertical accuracy. From this same LAZ, the observed 1 m terrain raster",
)

doc.save(OUTPUT)
check = Document(OUTPUT)
assert "0.196717 m/pixel" in check.paragraphs[46].text
assert "The 1 m DSM" in check.paragraphs[50].text
assert "0.5 m samples of the 1 m DSM" in check.paragraphs[52].text
print(OUTPUT)
