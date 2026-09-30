"""Synchronize the manuscript's cloud-derived claims with the definitive LAZ.

The source DOCX is read but never overwritten. Other geophysical claims and
figures are preserved; only point-cloud measurements and Figures 2/5 change.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches


OLD = {
    "2,779,817": "2,810,862",
    "20.41%": "20.64%",
    "44,091": "44,429",
    "46.77%": "47.13%",
    "53.23%": "52.87%",
    "7.37 points": "7.45 points",
    "Only 11 profiles": "Only 12 profiles",
    "remaining 21": "remaining 20",
    "Only 11 of the 32": "Only 12 of the 32",
}

AB_LIST_OLD = "P1, P2, P5, P6, P7, P9, P10, P11, P18, P23, and P24"
AB_LIST_NEW = "P1, P2, P3, P5, P6, P7, P9, P10, P11, P18, P23, and P24"

FIG2_CAPTION = (
    "Figure 2. Spatial context and final-cloud profile checks (WGS 84 / UTM zone 18S). "
    "(a) A 1 m cartographic terrain display based on class-2 Ground cell medians. "
    "Missing Ground cells are interpolated solely for visual context; the separately "
    "supplied observed-cell raster retains NoData in those locations. White lines "
    "show the 32 nominally west-to-east, 180 m profiles. The red line joins their "
    "mapped 60 m intersection positions as a local reference, not a demonstrated "
    "fault trace. (b-f) P32, P21, P18, P11, and P4, north to south, from the same "
    "definitive LAZ. Gray and blue dots are raw non-Ground and Ground observations "
    "within ±1 m of each line; class-7 noise is excluded. Orange dashed profiles "
    "sample the 1 m DSM cell maxima of all non-noise classes at 0.5 m intervals. "
    "Black segments represent directly supported class-2 Ground medians in "
    "0.5 m longitudinal bins of ±1 m swaths; unsupported bins remain empty. "
    "Gray dashed curves are an interpolated Ground-only context, not direct "
    "observations or inputs to morphometry. The red dotted line indicates the "
    "mapped lineament position at 60 m. Panels use common 0–180 m and 595–645 m "
    "axes at approximately 1:1 horizontal-to-vertical data scale. P32, P21, and "
    "P4 are class C because direct Ground support is insufficient; they are shown "
    "for quality control, not as independent scarp measurements."
)


def replace_picture(paragraph, path: Path, *, height: float,
                    page_break_before: bool = False) -> None:
    for run in paragraph.runs:
        run.clear()
    paragraph.add_run().add_picture(str(path), height=Inches(height))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = page_break_before


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--map", type=Path, required=True)
    parser.add_argument("--profiles", type=Path, required=True)
    parser.add_argument("--density", type=Path, required=True)
    args = parser.parse_args()
    doc = Document(args.source)
    pars = doc.paragraphs

    for p in pars:
        for run in p.runs:
            if not run.text:
                continue
            text = run.text
            for old, new in OLD.items():
                text = text.replace(old, new)
            text = text.replace(AB_LIST_OLD, AB_LIST_NEW)
            run.text = text

    pars[83].runs[0].text = (
        "Raw-Ground swath screening classified seven profiles as A (high direct "
        "support), five as B (moderate support), and 20 as C (insufficient support). "
        f"The A/B profiles are {AB_LIST_NEW}. Their Ground coverage in the "
        "morphometric search zone ranges from 68.6% to 97.1%; coverage near the "
        "mapped-line intersection ranges from 62.5% to 100%. The full support "
        "metrics, explicit A/B/C criteria, and all 32 profile bins are supplied "
        "with the reproducible quality-control files. Figure 2b-f provides "
        "geographically distributed raw-cloud examples and distinguishes direct "
        "Ground observations from connections across unsupported gaps."
    )
    pars[78].runs[0].text = (
        "This section reports the spatial distribution of direct Ground support "
        "and the resulting limits on terrain-profile interpretation, followed by "
        "the seismic-refraction-tomography results and their inversion-stability "
        "assessment. Interpolated DEM features are used for spatial context, not "
        "as independent scarp-height or displacement measurements."
    )
    pars[45].runs[0].text = FIG2_CAPTION
    replace_picture(pars[44], args.map, height=5.9)
    new_element = OxmlElement("w:p")
    pars[44]._p.addnext(new_element)
    profile_paragraph = Paragraph(new_element, pars[44]._parent)
    profile_paragraph.add_run().add_picture(str(args.profiles), height=Inches(6.80))
    profile_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    profile_paragraph.paragraph_format.page_break_before = True
    replace_picture(pars[81], args.density, height=6.25)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
