"""Apply final-cloud photogrammetry edits to the author's latest marked DOCX.

Every red-underlined source run is preserved verbatim and in order. The source
file is never overwritten. Revised photogrammetry text receives yellow Word
highlighting so that the author can review it separately from red SRT markup.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.shared import Inches


def is_red(run: Run) -> bool:
    color = run.font.color.rgb
    return bool(run.font.underline) and color is not None and str(color) == "C00000"


def red_runs(doc) -> list[tuple[str, bool, str]]:
    return [(r.text, bool(r.font.underline), str(r.font.color.rgb))
            for p in doc.paragraphs for r in p.runs if r.text and is_red(r)]


def yellow(run: Run) -> None:
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def replace_plain(paragraph: Paragraph, text: str) -> None:
    if any(is_red(r) and r.text for r in paragraph.runs):
        raise ValueError("Refusing to replace a paragraph containing red SRT markup")
    if not paragraph.runs:
        paragraph.add_run(text)
        yellow(paragraph.runs[0])
        return
    paragraph.runs[0].text = text
    yellow(paragraph.runs[0])
    for run in paragraph.runs[1:]:
        run.text = ""


def set_yellow_run(paragraph: Paragraph, index: int, text: str) -> None:
    run = paragraph.runs[index]
    if is_red(run):
        raise ValueError("Refusing to change a red SRT run")
    run.text = text
    yellow(run)


def add_after(paragraph: Paragraph, source_run: Run, text: str) -> Run:
    element = OxmlElement("w:r")
    source_run._r.addnext(element)
    new_run = Run(element, paragraph)
    new_run.text = text
    return new_run


def add_before(paragraph: Paragraph, source_run: Run, text: str) -> Run:
    element = OxmlElement("w:r")
    source_run._r.addprevious(element)
    new_run = Run(element, paragraph)
    new_run.text = text
    return new_run


def replace_picture(paragraph: Paragraph, path: Path, height: float) -> None:
    for run in paragraph.runs:
        run.clear()
    paragraph.add_run().add_picture(str(path), height=Inches(height))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


FIG2 = (
    "Final-cloud spatial context and point-cloud profile checks (WGS 84 / UTM zone 18S). "
    "(a) A 1 m terrain display based on class-2 Ground cell medians. Gaps are interpolated "
    "only for visual context, not interpreted as direct terrain observations; the observed-cell "
    "raster retains NoData. Hillshade was generated with azimuth 315°, illumination altitude "
    "45°, and vertical exaggeration 1. White lines show the 32 nominally west-to-east, 180 m "
    "profiles. The red line joins their mapped 60 m intersections as a local reference, not "
    "a demonstrated fault trace. (b–f) P32, P21, P18, P11, and P4, north to south. Gray and "
    "blue points are raw non-Ground and class-2 Ground observations within ±1 m of each line; "
    "class-7 noise is excluded. The orange dashed line samples 1 m DSM cell maxima from all "
    "non-noise classes at 0.5 m intervals. Black segments represent directly observed Ground "
    "medians in 0.5 m longitudinal bins of ±1 m swaths; empty bins are not connected. Gray "
    "dashed curves show interpolated Ground-only context. The red dotted line marks 60 m. "
    "The panels share 0–180 m and 595–645 m axes at approximately 1:1 metric scaling. "
    "P32, P21, and P4 lack sufficient direct Ground support for independent scarp measurements."
)

FIG5 = (
    "Quality control of the definitive UAV point cloud in 2 × 2 m cells (EPSG:32718). "
    "(A) Density of all reconstructed points; (B) class-2 Ground density; (C) Ground retention "
    "as a percentage of all points; and (D) direct Ground support. A and B use the same "
    "display scale capped at the 99th percentile of all-point density (105.0 points m⁻²). "
    "The LAZ contains 13,619,241 points, including 2,810,862 Ground (20.64%). Densities "
    "over the common occupied footprint are 36.12 and 7.45 points m⁻², respectively. "
    "Ground occurs in 44,429 of 94,274 occupied cells (47.13%); the others are not direct "
    "bare-earth observations."
)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--map", type=Path, required=True)
    parser.add_argument("--profiles", type=Path, required=True)
    parser.add_argument("--density", type=Path, required=True)
    args = parser.parse_args()
    doc = Document(args.source)
    pars = doc.paragraphs
    source_red = red_runs(doc)

    replace_plain(pars[0],
        "Direct Ground Support and Shallow P-wave Velocity Heterogeneity at a Candidate "
        "Southern Continuation of the Porotoyacu Fault, Ecuadorian Amazon")
    replace_plain(pars[21],
        "Keywords: UAV Structure-from-Motion; Ground support; classified point cloud; "
        "Seismic refraction tomography; Porotoyacu Fault")

    # Abstract: retain the existing red SRT run sequence and its adjacent
    # non-red data-fit language; replace only photo-derived assertions.
    set_yellow_run(pars[20], 0,
        "This study evaluates direct terrain support and shallow P-wave velocity "
        "heterogeneity at a candidate southern continuation of the Porotoyacu Fault. "
        "The definitive UAV-SfM LAZ contains 13,619,241 classified points, including "
        "2,810,862 class-2 Ground points (20.64%). Ground occurs in 47.13% of occupied "
        "2 × 2 m cells. Thirty-two 180 m profiles were screened with raw Ground points in "
        "±5 m swaths and 1 m bins; 12 have high or moderate direct support, whereas 20 "
        "cannot support independent morphometric measurement. ")
    add_after(pars[20], pars[20].runs[0],
        "A single 69 m SW–NE SRT profile was also evaluated. ")
    # Re-index after inserting the plain SRT run.
    photo_assertion = next(r for r in pars[20].runs
                           if r.text.startswith("Topographic analysis documents"))
    photo_assertion.text = (
        "The observed terrain data show a local elevation transition but do not "
        "verify continuous fault-trace extension, vertical separation, or a scarp height. The "
    )
    yellow(photo_assertion)
    final_run = pars[20].runs[-1]
    tail = "The confirmed southward extension of the fault trace"
    if tail not in final_run.text:
        raise ValueError("Expected unsupported abstract conclusion not found")
    final_run.text = final_run.text.split(tail)[0]
    ending = pars[20].add_run(
        "Together, the terrain and velocity observations motivate further field testing; "
        "they do not establish a fault length, present activity, or a seismic-hazard parameter."
    )
    yellow(ending)

    set_yellow_run(pars[26], 0,
        "UAV Structure-from-Motion can provide detailed surface reconstructions, but "
        "raster cell spacing is not vertical accuracy and a bare-earth model depends on "
        "where Ground-classified points actually occur. This distinction is especially "
        "important beneath tropical vegetation and on roofs or other built surfaces. ")
    add_after(pars[26], pars[26].runs[0],
        "Shallow seismic methods can add complementary subsurface information. ")
    replace_plain(pars[28],
        "The Porotoyacu Fault is mapped regionally as part of the Napo-Cutucú fault "
        "system (Figure 1C). Regional mapping motivates investigation of a possible "
        "southern continuation, but the local surface continuity and origin of nearby "
        "topographic breaks have not been independently verified. This study therefore "
        "treats the surface feature as a candidate lineament, not proof of active rupture.")
    set_yellow_run(pars[29], 0,
        "To address this gap, the study combines final-cloud Ground quality control "
        "with one SRT line. Its objectives are to (1) quantify direct Ground support "
        "along 32 local profiles; and (2) describe ")
    set_yellow_run(pars[29], 4,
        "through tomographic inversion; and (3) examine whether the local surface "
        "and velocity observations justify a testable structural hypothesis, without "
        "claiming verified fault continuity, displacement, activity, or hazard parameters.")
    set_yellow_run(pars[41], 0,
        "We used the definitive classified UAV point cloud to quantify direct "
        "terrain support and one SRT line to describe shallow subsurface variability. "
        "The interpolated terrain display was not treated as direct topographic "
        "measurement, while subsurface ")
    replace_plain(pars[43],
        "The study area lies in the Ecuadorian Sub-Andean Zone near the mapped "
        "southern Porotoyacu Fault segment (Figure 2). The UAV survey covered a local "
        "topographic transition near latitude −1.016118°, longitude −77.805551° "
        "(WGS 84). The feature is investigated as a candidate surface lineament rather "
        "than assumed to be a fault scarp or continuous along strike.")

    replace_picture(pars[44], args.map, 5.9)
    new_element = OxmlElement("w:p")
    pars[44]._p.addnext(new_element)
    profile_paragraph = Paragraph(new_element, pars[44]._parent)
    profile_paragraph.add_run().add_picture(str(args.profiles), height=Inches(6.8))
    profile_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    profile_paragraph.paragraph_format.page_break_before = True
    set_yellow_run(pars[45], 0, "Figure 2. ")
    set_yellow_run(pars[45], 1, FIG2)

    replace_plain(pars[49],
        "An initial dense reconstruction was classified in Agisoft Metashape "
        "Professional 2.0.0 using the documented slope settings (maximum angle "
        "7°, maximum distance 0.1 m, cell size 25 m). Point-cloud quality control "
        "subsequently produced the definitive classified LAZ used here. From that "
        "same LAZ, the observed 1 m terrain raster stores the median elevation of "
        "class-2 Ground points per cell and retains NoData in empty cells. The 1 m "
        "DSM stores the cell maximum of all classes except class-7 noise, including "
        "class-11 Road Surface. Interpolated terrain cells are used only as visual "
        "context, not as directly observed ground or evidence of vertical accuracy. "
        "All rasters are in EPSG:32718.")
    replace_plain(pars[50],
        "For quantitative validation, LAS classifications were preserved from the "
        "definitive LAZ. All-point and class-2 Ground counts were evaluated on one "
        "2 × 2 m grid. The common footprint consists of cells containing at least "
        "one reconstructed point; counts divided by 4 m² give local density. "
        "Ground retention is 100 × Nground/Nall, and direct terrain support means "
        "at least one Ground point in a cell. All other LAS classes are excluded "
        "from the observed DTM, but are not all interpreted as vegetation. Class-11 "
        "Road Surface contributes to the DSM, not to the Ground-only DTM.")
    replace_plain(pars[51],
        "Thirty-two nominally transverse 180 m profiles were retained at their "
        "mapped west-to-east positions (P1 south to P32 north; azimuths 66.11°–94.52°). "
        "Each crosses the local mapped lineament position at 60 m. Class-2 points "
        "were extracted in ±5 m swaths, summarized as medians and interquartile "
        "ranges in 1 m longitudinal bins, and considered directly supported only "
        "when a bin contained at least three Ground points; empty bins were not "
        "filled. ±2.5 m and ±7.5 m swaths were retained for sensitivity checks. "
        "The five illustrative plots also show raw points in ±1 m corridors, "
        "direct 0.5 m Ground medians, and 0.5 m samples of the 1 m DSM and "
        "interpolated DTM context. The context curves are not morphometric data.")
    srt_method = pars[53].runs[0]
    old_lead = (
        "To characterize the shallow subsurface P-wave velocity structure beneath "
        "the mapped scarp, a 69 m long seismic refraction tomography profile was "
        "acquired along the seismic transect shown in Figure 2. "
    )
    if not srt_method.text.startswith(old_lead):
        raise ValueError("SRT method cross-reference changed in source")
    srt_method.text = srt_method.text[len(old_lead):]
    yellow(add_before(pars[53], srt_method,
        "To characterize local shallow P-wave velocity heterogeneity, a 69 m "
        "seismic refraction tomography profile was acquired in the study area. "))

    set_yellow_run(pars[77], 0,
        "This section first reports direct Ground support and the limits it "
        "places on terrain-profile interpretation. The SRT results then follow, "
        "including inversion stability and the ")
    set_yellow_run(pars[80], 0,
        "The definitive classified cloud contains 13,619,241 points and "
        "2,810,862 class-2 Ground points (20.64%). Across 94,274 occupied "
        "2 × 2 m cells (0.377096 km²), all-point density is 36.12 points m⁻² "
        "and Ground density is 7.45 points m⁻². Direct Ground support occurs "
        "in 44,429 cells (47.13%); the other 52.87% contain reconstructed "
        "points but no direct terrain observation (Figure ")
    replace_picture(pars[81], args.density, 6.25)
    set_yellow_run(pars[82], 2, FIG5)
    replace_plain(pars[83],
        "The 32 profiles were screened using raw Ground observations rather than "
        "treating continuous DTM raster coverage as equivalent to direct terrain "
        "measurement. The ±5 m, 1 m-bin test classified seven profiles as A "
        "(high direct support), five as B (moderate), and 20 as C (insufficient). "
        "The A/B profiles are P1, P2, P3, P5, P6, P7, P9, P10, P11, P18, P23, "
        "and P24. Their Ground coverage in the 30–170 m search zone spans "
        "68.6%–97.1%; near the mapped intersection (40–80 m), it spans "
        "62.5%–100%. Full criteria and bins are provided with the profile QC files.")
    replace_plain(pars[84],
        "A/B designation supports only local description of directly observed "
        "terrain. It does not validate a base, crest, scarp height, tectonic "
        "vertical separation, or displacement because independent far-field "
        "reference surfaces and an uncertainty model are absent. No mean height "
        "or two-domain morphometric partition is calculated from these profiles.")
    replace_plain(pars[85],
        "The representative plots contrast the DSM with sparse or continuous "
        "Ground observations. P11 has high direct support and shows a local "
        "west-to-east elevation rise; P18 has moderate support. P4, P21, and "
        "P32 have insufficient support in relevant intervals, even where an "
        "interpolated DTM appears smooth. DSM–DTM differences may reflect "
        "vegetation, buildings, other above-ground objects, or Ground gaps; "
        "they are not direct measurements of fault relief.")
    replace_plain(pars[87],
        "Direct support varies markedly along the 32 lines. The 20 class-C "
        "profiles, including P21, P31, and P32, are excluded from independent "
        "morphometric interpretation. A continuous raster can depict a local "
        "transition through those gaps, but it does not establish a continuous "
        "geomorphic or tectonic feature along strike.")

    replace_plain(pars[117],
        "The combined observations allow a local, support-qualified description "
        "of terrain and a separate description of the shallow SRT velocity field. "
        "Their spatial relationship is evaluated as a working hypothesis, not "
        "as proof of fault continuity, geometry, or activity.")
    replace_plain(pars[118], "5.1. Direct-Ground Support and Local Surface Morphology")
    replace_plain(pars[119],
        "Only 12 of the 32 nominally transverse profiles satisfy the A/B direct-" 
        "Ground support screen. They provide local observations of elevation "
        "transitions, not a complete along-strike morphometric sample. They do "
        "not establish a continuous scarp or a mean tectonic separation.")
    replace_plain(pars[120],
        "The remaining 20 profiles have substantial gaps within the search "
        "zone or near the mapped-line window. Vegetation, built surfaces, and "
        "sparse Ground observations can cause terrain rasters to interpolate "
        "across these areas. No apparent depression, step, or relief in a "
        "class-C profile is used as an independent tectonic measurement.")
    replace_plain(pars[121],
        "Consequently, the previously proposed 700 m southward continuation, "
        "115 m positional refinement, average relative height, and two-sector "
        "morphometric partition are not retained as demonstrated results. Testing "
        "them would require independent field mapping, support-complete profiles, "
        "far-field reference surfaces, and an uncertainty assessment.")
    set_yellow_run(pars[127], 0,
        "The directly supported Ground points document a local topographic "
        "transition, but do not identify a unique tectonic origin or a continuous "
        "surface trace. Differential erosion and anthropogenic modification "
        "remain alternative explanations. ")
    replace_plain(pars[134], "5.5. Limits on Fault Length and Hazard Inference")
    replace_plain(pars[135],
        "The available Ground observations do not verify a 700 m extension "
        "of the mapped fault, establish a single seismogenic segment, or "
        "supply a source geometry for hazard modeling. The local lineament "
        "remains a target for field verification.")
    replace_plain(pars[139],
        "No independent Ground Control Points, checkpoints, or Total Station "
        "calibration were available to quantify absolute vertical accuracy. "
        "Cell size and RTK georeferencing do not substitute for that test. "
        "Furthermore, variable Ground support and anthropogenic modification "
        "limit profile interpretation; apparent relief cannot be converted "
        "to a minimum tectonic displacement without an independent reference.")
    set_yellow_run(pars[140], 0,
        "Structure-from-Motion photogrammetry cannot penetrate dense canopy "
        "as airborne LiDAR can. Ground gaps and possible residual vegetation "
        "therefore limit interpretation in the northern sector rather than "
        "merely adding a small error to a known scarp height. No geomorphic "
        "marker age or paleoseismic constraint was available to infer recent "
        "activity or a slip rate. ")
    set_yellow_run(pars[142], 0,
        "This study combines final-cloud UAV photogrammetric quality control "
        "with shallow seismic refraction tomography ")
    set_yellow_run(pars[142], 2,
        " The final cloud has 13,619,241 points, including 2,810,862 class-2 "
        "Ground points (20.64%); direct Ground occurs in 47.13% of occupied "
        "2 × 2 m cells. Twelve of 32 profiles meet A/B support criteria. They "
        "record local elevation transitions but do not independently verify a "
        "continuous 700 m scarp, vertical separation, or displacement.")

    if red_runs(doc) != source_red:
        raise AssertionError("Red SRT markup changed; output not saved")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    # Round-trip verification: python-docx sometimes normalizes run XML.
    if red_runs(Document(args.output)) != source_red:
        raise AssertionError("Red SRT markup changed during DOCX save")
    print(f"Saved {args.output}; {len(source_red)} red SRT runs preserved")


if __name__ == "__main__":
    main()
