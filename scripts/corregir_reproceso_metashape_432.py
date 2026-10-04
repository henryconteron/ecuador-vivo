"""Correct the documented 432-camera Metashape workflow in the latest manuscript.

The source is the already unified, yellow-highlighted manuscript. This script
changes only two photogrammetry methods paragraphs and one limitations paragraph,
preserving the user's red-underlined tomography revisions and embedded figures.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_DIR = ROOT / "output" / "porotoyacu_nube_definitiva_20260923" / "manuscrito"
SOURCE = MANUSCRIPT_DIR / "Manuscript_Conteron_etal_2026_UNIFICADO_FOTOGRAMETRIA_AMARILLO.docx"
OUTPUT = MANUSCRIPT_DIR / "Manuscript_Conteron_etal_2026_FOTOGRAMETRIA_432_VERIFICADA.docx"


REPLACEMENTS = {
    49: (
        "We conducted an aerial survey using a DJI Phantom 4 RTK drone, "
        "acquiring 432 images over approximately 0.36 km². Camera positions "
        "from the onboard RTK system provided direct georeferencing, but no "
        "independent ground-control checkpoints were available. Images were "
        "processed in Agisoft Metashape Professional 2.0.0. Alignment used high "
        "accuracy, generic preselection, source-reference preselection, a "
        "40,000 key-point limit, a 10,000 tie-point limit, and adaptive "
        "camera-model fitting. All 432 cameras aligned; the internal tie-point "
        "reprojection error was 0.748 pixels. Depth maps were generated for "
        "all 432 cameras at low quality with moderate filtering and a maximum "
        "of 16 neighbors. The initial dense point cloud contained 13,619,251 "
        "points."
    ),
    50: (
        "The initial dense cloud was classified in Agisoft Metashape "
        "Professional 2.0.0 using a maximum terrain angle of 7°, maximum "
        "distance of 0.1 m, and cell size of 25 m. At this initial stage, "
        "2,840,463 points were labeled Ground. Subsequent manual point-level "
        "review and reclassification yielded the definitive LAZ analyzed here: "
        "13,619,241 total points, including 2,810,862 class-2 Ground points. "
        "The final LAZ, not the initial Metashape classification, underpins "
        "the profile and density analyses. From this same LAZ, the observed "
        "1 m terrain raster stores the median elevation of class-2 Ground "
        "points per cell and retains NoData in empty cells. The 1 m DSM "
        "stores the cell maximum of all classes except class-7 noise, "
        "including class-11 Road Surface. Interpolated terrain cells are used "
        "only as visual context, not as directly observed ground or evidence "
        "of vertical accuracy. All rasters are in EPSG:32718."
    ),
    140: (
        "No independent Ground Control Points, checkpoints, or Total Station "
        "calibration were available to quantify absolute vertical accuracy. "
        "The 0.748-pixel reprojection error describes internal image alignment, "
        "not absolute elevation accuracy; neither cell size nor RTK "
        "georeferencing substitutes for an independent vertical check. "
        "Furthermore, variable Ground support and anthropogenic modification "
        "limit profile interpretation; apparent relief cannot be converted "
        "to a minimum tectonic displacement without an independent reference."
    ),
}


def red_segments(document):
    segments = []
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            rgb = run.font.color.rgb if run.font.color else None
            if rgb is not None and str(rgb).upper() in {"C00000", "FF0000"} and run.underline:
                segments.append((run.text, str(rgb).upper(), run.underline))
    return segments


def main():
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing deliverable: {OUTPUT}")

    document = Document(SOURCE)
    original_red = red_segments(document)
    original_figures = len(document.inline_shapes)
    original_paragraphs = len(document.paragraphs)

    for index, replacement in REPLACEMENTS.items():
        paragraph = document.paragraphs[index]
        if len(paragraph.runs) != 1:
            raise AssertionError(f"Expected one text run in paragraph {index}")
        if index == 49 and ("431 were successfully aligned" not in paragraph.text or "430 contributed depth maps" not in paragraph.text):
            raise AssertionError("The alignment paragraph no longer matches the expected draft")
        if index == 50 and "cell size 25 m" not in paragraph.text:
            raise AssertionError("The classification paragraph no longer matches the expected draft")
        if index == 140 and "No independent Ground Control Points" not in paragraph.text:
            raise AssertionError("The limitations paragraph no longer matches the expected draft")
        paragraph.runs[0].text = replacement
        paragraph.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW

    if red_segments(document) != original_red:
        raise AssertionError("Red-underlined tomography changes were altered")
    if len(document.inline_shapes) != original_figures:
        raise AssertionError("An embedded figure was altered")
    if len(document.paragraphs) != original_paragraphs:
        raise AssertionError("The document structure changed")

    document.save(OUTPUT)
    reopened = Document(OUTPUT)
    if red_segments(reopened) != original_red:
        raise AssertionError("Red-underlined tomography changes were not retained on save")
    assert len(reopened.inline_shapes) == original_figures
    assert len(reopened.paragraphs) == original_paragraphs
    assert "431 were successfully aligned" not in "\n".join(p.text for p in reopened.paragraphs)
    assert "430 contributed depth maps" not in "\n".join(p.text for p in reopened.paragraphs)
    for index, replacement in REPLACEMENTS.items():
        assert reopened.paragraphs[index].text == replacement
        assert reopened.paragraphs[index].runs[0].font.highlight_color == WD_COLOR_INDEX.YELLOW

    print(f"Saved {OUTPUT}")
    print(f"Preserved {len(original_red)} red-underlined tomography segments and {original_figures} figures")


if __name__ == "__main__":
    main()
